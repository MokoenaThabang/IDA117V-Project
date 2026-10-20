"""Assemble the IDA117V submission folder and ZIP from the current repository.

This script copies existing files. It does not train a model, rewrite a
result, or edit the original Markdown, notebooks, or pipeline.
"""

from __future__ import annotations

import csv
import json
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "submission_package"
ZIP_PATH = ROOT / "Cape_Town_Service_Request_Resolution_Classifier_Submission.zip"
ZIP_ROOT = "Cape_Town_Service_Request_Resolution_Classifier"

MARKDOWN_DOCUMENTS = (
    ("README.md", "Documentation/Repository_README.docx"),
    ("reports/contribution_statement.md", "Documentation/Contribution_Statement.docx"),
    ("reports/team_contribution_template.md", "Documentation/Team_Contribution_Template.docx"),
    ("reports/presentation_outline.md", "Documentation/Presentation_Outline.docx"),
    ("reports/data_source.md", "Documentation/Dataset_Source.docx"),
    ("reports/demonstration_checklist.md", "Documentation/Demonstration_Checklist.docx"),
    ("reports/experiment_log.md", "Documentation/Experiment_Log.docx"),
    ("reports/phase2_feasibility.md", "Documentation/Phase_2_Feasibility.docx"),
    ("reports/phase3_eda_findings.md", "Documentation/Phase_3_EDA_Findings.docx"),
    ("reports/phase4_baselines.md", "Documentation/Phase_4_Baselines.docx"),
    ("reports/phase5_models.md", "Documentation/Phase_5_Models.docx"),
    ("reports/phase6_evaluation.md", "Documentation/Phase_6_Evaluation.docx"),
    ("reports/reproducibility_audit.md", "Documentation/Reproducibility_Audit.docx"),
    ("reports/submission_checklist.md", "Documentation/Submission_Checklist.docx"),
    ("reports/submission_readiness.md", "Documentation/Submission_Readiness.docx"),
    ("reports/stakeholder_interview_guide.md", "Documentation/Stakeholder_Interview_Guide.docx"),
    ("docs/data-dictionary.md", "Dataset/Earlier_Modelling_Dictionary.docx"),
    ("docs/dataset-card.md", "Dataset/Earlier_Dataset_Card.docx"),
    ("docs/concept-notes.md", "Documentation/Concept_Notes.docx"),
    ("data/README.md", "Dataset/Data_Folder_Notes.docx"),
    ("data/processed/README.md", "Dataset/Processed_Folder_Notes.docx"),
    ("notebooks/README.md", "Notebooks/Notebook_Folder_Notes.docx"),
    ("app/README.md", "Source_Code/app/App_Notes.docx"),
    ("evidence/README.md", "Presentation/Evidence_Notes.docx"),
    ("figures/interpretations.md", "Results/Earlier_Figure_Interpretations.docx"),
)


def _read_json(relative: str) -> dict:
    return json.loads((ROOT / relative).read_text(encoding="utf-8"))


def _read_csv(relative: str) -> list[dict]:
    with (ROOT / relative).open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def _ignore_junk(_directory: str, names: list[str]) -> set[str]:
    skipped = set()
    for name in names:
        if name in {"__pycache__", ".ipynb_checkpoints", ".pytest_cache"} or name.endswith(".pyc"):
            skipped.add(name)
    return skipped


def _copy_tree(source: Path, destination: Path) -> None:
    if destination.exists():
        shutil.rmtree(destination)
    shutil.copytree(source, destination, ignore=_ignore_junk)


def _copy_file(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)


def _set_run_font(run, name: str, size: Pt, bold: bool = False, italic: bool = False, colour=None) -> None:
    run.bold = bold
    run.italic = italic
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = size
    if colour is not None:
        run.font.color.rgb = colour


def _add_page_number(paragraph) -> None:
    run_before = paragraph.add_run("Page ")
    _set_run_font(run_before, "Calibri", Pt(9), colour=RGBColor(0x55, 0x55, 0x55))

    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    _set_run_font(run, "Calibri", Pt(9), colour=RGBColor(0x55, 0x55, 0x55))
    run._r.append(field_begin)
    run._r.append(instr)
    run._r.append(field_end)

    run_after = paragraph.add_run(" of ")
    _set_run_font(run_after, "Calibri", Pt(9), colour=RGBColor(0x55, 0x55, 0x55))

    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " NUMPAGES "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    _set_run_font(run, "Calibri", Pt(9), colour=RGBColor(0x55, 0x55, 0x55))
    run._r.append(field_begin)
    run._r.append(instr)
    run._r.append(field_end)


def _shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")
    properties.append(shading)


def _apply_document_defaults(document: Document, title: str) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    header = section.header.paragraphs[0]
    header.alignment = 1
    run = header.add_run("Cape Town service-request resolution classifier  ·  IDA117V")
    _set_run_font(run, "Calibri", Pt(9), colour=RGBColor(0x44, 0x55, 0x66))
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    _add_page_number(footer)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    paragraph_format = normal.paragraph_format
    paragraph_format.space_after = Pt(8)
    paragraph_format.space_before = Pt(0)
    paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph_format.line_spacing = 1.08

    sizes = {1: 18, 2: 14, 3: 12, 4: 11}
    for level, size in sizes.items():
        style = document.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x1B, 0x3A, 0x4B)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
    document.core_properties.title = title
    document.core_properties.subject = "IDA117V Cape Town service-request resolution classifier"
    document.core_properties.category = "Academic submission package"


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "22")
    font = OxmlElement("w:rFonts")
    font.set(qn("w:ascii"), "Calibri")
    font.set(qn("w:hAnsi"), "Calibri")
    properties.extend([font, colour, underline, size])
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.set(qn("xml:space"), "preserve")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE = re.compile(
    r"(\*\*[^*]+?\*\*|`[^`]+`|\[[^\]\n]+\]\([^)\n]+\)|(?<!\*)\*[^*\n]+?\*(?!\*))"
)


def _add_inline(paragraph, text: str, size: Pt = Pt(11), base_font: str = "Calibri") -> None:
    position = 0
    for match in INLINE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            _set_run_font(run, base_font, size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            _set_run_font(run, base_font, size, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, "Consolas", Pt(10))
        elif token.startswith("["):
            label, url = token[1:-1].split("](", 1)
            _add_hyperlink(paragraph, label, url)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, base_font, size, italic=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        _set_run_font(run, base_font, size)


def _add_code_block(document: Document, lines: list[str]) -> None:
    table = document.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    _shade_cell(cell, "F4F6F8")
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(2)
    content = "\n".join(lines) if lines else ""
    run = paragraph.add_run(content)
    _set_run_font(run, "Consolas", Pt(9))
    document.add_paragraph()


def _add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(rows):
        for column_index in range(width):
            value = row[column_index] if column_index < len(row) else ""
            cell = table.cell(row_index, column_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.space_before = Pt(2)
            _add_inline(paragraph, value.replace("<br>", " "), size=Pt(9))
            if row_index == 0:
                _shade_cell(cell, "E6EEF2")
                for run in paragraph.runs:
                    run.bold = True
    document.add_paragraph()


def _split_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def _is_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?", line.strip()))


def convert_markdown(source: Path, destination: Path, preface: list[str] | None = None) -> None:
    document = Document()
    _apply_document_defaults(document, source.stem.replace("_", " "))
    if preface:
        for line in preface:
            paragraph = document.add_paragraph()
            _add_inline(paragraph, line)
            for run in paragraph.runs:
                run.italic = True
    text = source.read_text(encoding="utf-8")
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            code: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            _add_code_block(document, code)
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and _is_separator(lines[index + 1]):
            rows = [_split_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_split_row(lines[index].strip()))
                index += 1
            _add_table(document, rows)
            continue
        image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image:
            caption, raw_path = image.group(1), image.group(2)
            image_path = (source.parent / raw_path).resolve()
            paragraph = document.add_paragraph()
            if image_path.is_file():
                run = paragraph.add_run()
                run.add_picture(str(image_path), width=Inches(6.2))
                if caption:
                    caption_paragraph = document.add_paragraph()
                    run = caption_paragraph.add_run(caption)
                    _set_run_font(run, "Calibri", Pt(9), italic=True, colour=RGBColor(0x44, 0x55, 0x66))
            else:
                _add_inline(paragraph, f"Image not found at `{raw_path}`. {caption}")
            index += 1
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            document.add_heading(heading.group(2), level=len(heading.group(1)))
            index += 1
            continue
        if re.fullmatch(r"(-{3,}|\*{3,})", stripped):
            index += 1
            continue
        bullet = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline(paragraph, bullet.group(2))
            index += 1
            continue
        number = re.match(r"^(\s*)\d+[.)]\s+(.*)$", line)
        if number:
            paragraph = document.add_paragraph(style="List Number")
            _add_inline(paragraph, number.group(2))
            index += 1
            continue
        quote = re.match(r"^>\s?(.*)$", stripped)
        if quote:
            paragraph = document.add_paragraph()
            _add_inline(paragraph, quote.group(1))
            for run in paragraph.runs:
                run.italic = True
            index += 1
            continue
        if stripped == "":
            index += 1
            continue
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines) and lines[index].strip() and not _starts_block(lines[index]):
            paragraph_lines.append(lines[index].strip())
            index += 1
        paragraph = document.add_paragraph()
        _add_inline(paragraph, " ".join(paragraph_lines))
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def _starts_block(line: str) -> bool:
    stripped = line.strip()
    return bool(
        stripped.startswith("```")
        or stripped.startswith("|")
        or stripped.startswith("#")
        or stripped.startswith(">")
        or stripped.startswith("!")
        or re.match(r"^[-*]\s+", stripped)
        or re.match(r"^\d+[.)]\s+", stripped)
        or re.fullmatch(r"(-{3,}|\*{3,})", stripped)
    )


def _paragraph(document: Document, text: str, *, bullet: bool = False, number: bool = False) -> None:
    style = "List Bullet" if bullet else "List Number" if number else None
    paragraph = document.add_paragraph(style=style) if style else document.add_paragraph()
    _add_inline(paragraph, text)


def _code(document: Document, text: str) -> None:
    _add_code_block(document, text.splitlines())


def write_project_overview(destination: Path) -> None:
    decision = _read_json("results/tables/06_frozen_decision.json")
    question = _read_json("results/tables/13_question62_record.json")
    logistic = _read_json("results/tables/11_logistic_convergence.json")
    document = Document()
    _apply_document_defaults(document, "Project overview")
    document.add_heading("Project overview", 1)
    _paragraph(
        document,
        "This overview is assembled from the repository README, the saved decision files, and the technical report. It does not replace `Technical_Report.pdf`.",
    )
    document.add_heading("Purpose", 2)
    _paragraph(
        document,
        "At the moment a City of Cape Town municipal service request is logged, estimate whether it is likely to be Resolved Fast, Resolved Slow, or Unresolved. The estimate is a review aid for a person in the municipal service department. It is not a service decision.",
    )
    document.add_heading("Label and split", 2)
    _paragraph(
        document,
        "Resolved Fast means completed within seven calendar days. That cut is our project rule, not a City standard. The split id recorded for the frozen decision is `"
        + decision["split_id"]
        + "`. Training ends on 31 December 2024, validation is 2025, and the locked test is 1 January 2026 through 30 April 2026. Seed 42.",
    )
    document.add_heading("Champion", 2)
    _paragraph(
        document,
        "The champion was selected on validation macro-F1. The locked test was not used to select or tune it. The saved decision names `"
        + decision["champion_family"]
        + "`, selected `"
        + decision["selected"]
        + "`, with max_iter "
        + str(decision["max_iter"])
        + ". Phase 5 validation macro-F1 is "
        + str(decision["phase5_champion_validation_macro_f1"])
        + ". The frozen untuned validation macro-F1 is "
        + str(decision["untuned_validation_macro_f1"])
        + ", which is higher than the tuned validation macro-F1 of "
        + str(decision["tuned_validation_macro_f1"])
        + ". Locked-test macro-F1 is "
        + str(decision["locked_test_macro_f1"])
        + ".",
    )
    _paragraph(
        document,
        "Features recorded in the frozen decision: "
        + ", ".join(f"`{name}`" for name in decision["features"])
        + ". Work centre is excluded because the extract has no assignment time.",
    )
    document.add_heading("Question 6.2", 2)
    _paragraph(
        document,
        "The Question 6.2 comparator is the embedding network, not the champion. The saved record compares `"
        + question["champion"]
        + "` with `"
        + question["comparator"]
        + "` on "
        + f"{question['test_rows']:,}"
        + " locked-test rows. The continuity-corrected statistic is "
        + str(question["statistic"])
        + " and the p-value is "
        + str(question["p_value"])
        + ". `champion_changed` is "
        + str(question["champion_changed"])
        + ". The historical multilayer-perceptron McNemar table remains a separate result.",
    )
    document.add_heading("Logistic regression", 2)
    _paragraph(
        document,
        "The recorded 200-iteration logistic regression did not converge (`converged` is "
        + str(logistic["converged"]).lower()
        + ", `convergence_warning` is "
        + str(logistic["convergence_warning"]).lower()
        + "). Solver `"
        + logistic["solver"]
        + "`, max_iter "
        + str(logistic["historical_max_iter"])
        + ", validation macro-F1 "
        + str(logistic["recorded_200_iteration_validation_macro_f1"])
        + ". A 1000-iteration check scored "
        + str(logistic["this_run_validation_macro_f1"])
        + " and was not adopted.",
    )
    document.add_heading("Safeguards", 2)
    _paragraph(
        document,
        "Predictions must support human review. They must not be used to deny or reduce service to an area. This is an academic prototype for an IDA117V project. It is not an official City of Cape Town system. The POPIA section of the report describes the fields that were withheld. It does not certify compliance.",
    )
    document.add_heading("Cover fields the team must still supply", 2)
    for item in (
        "The report cover names Mpumelelo Ngozo (218409202), Selowa Tumiso (218340172), Ualwela Netshidzati (220069117), and Thabang Mokoena (240668823).",
        "Group number: not recorded. Complete manually.",
        "Lecturer: not recorded. Complete manually.",
        "Submission date: not recorded. Complete manually.",
        "Contribution categories name one member each. Signature cells contain that member’s initials. Upper rows are dated 5–8 October 2026. The lower half is dated 12–18 October 2026.",
        "Presentation attendance and defence outcome: not in the repository.",
    ):
        _paragraph(document, item, bullet=True)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def write_readme(destination: Path) -> None:
    decision = _read_json("results/tables/06_frozen_decision.json")
    question = _read_json("results/tables/13_question62_record.json")
    logistic = _read_json("results/tables/11_logistic_convergence.json")
    locked = _read_csv("results/tables/06_locked_test_summary.csv")[0]
    per_class = _read_csv("results/tables/06_locked_test_per_class.csv")
    slow = next(row for row in per_class if row["label"] == "Resolved Slow")
    document = Document()
    _apply_document_defaults(document, "README")
    document.add_heading("Cape Town service-request resolution classifier", 1)
    _paragraph(
        document,
        "IDA117V academic project. This README describes the repository that produced the files in this package. Paths such as `streamlit run app/app.py` are the commands for that repository layout. Inside this package, the same files are under `Source_Code/`, `Notebooks/`, `Model_Artifacts/`, and `Results/`.",
    )

    document.add_heading("1. Project title and purpose", 2)
    _paragraph(
        document,
        "Our project estimates, at logging time, which of three resolution outcomes is likely for a City of Cape Town SAP C3 service request: Resolved Fast, Resolved Slow, or Unresolved.",
    )
    document.add_heading("2. Problem", 2)
    _paragraph(
        document,
        "A reviewer needs an early indication of which newly logged requests may need attention. Accuracy alone would favour the majority class, Resolved Fast. The primary technical metric is macro-F1.",
    )
    document.add_heading("3. Decision-maker and appropriate use", 2)
    _paragraph(
        document,
        "The decision-maker for our estimate is a reviewer in the municipal service department that receives the request. The estimate is a queue aid. A person remains between the score and any operational action. Predictions must support human review. They must not be used to deny or reduce service to an area. We treat this as an academic prototype and not as an official City of Cape Town system.",
    )
    document.add_heading("4. Dataset", 2)
    _paragraph(
        document,
        "The source is the City of Cape Town Open Data Portal feature layer titled Service Requests 2023 until 06 October 2026 (Deprecated). The project plan names it SAP C3 Notifications. The dataset page is https://odp-cctegis.opendata.arcgis.com/datasets/90711027d33940b5a06e96ad8a8f7ede_0/explore. The licence is the City of Cape Town Open Data terms of use. It is custom, not a Creative Commons licence. The portal count recorded in `reports/data_source.md` is 5,372,466 rows, with created-on dates from 1 January 2023 through 6 October 2026. The repository code is MIT and does not cover the City data.",
    )
    document.add_heading("5. Directory structure", 2)
    _paragraph(document, "Repository layout used by the notebooks and the application:", bullet=False)
    for item in (
        "`app/` — Streamlit interface. It loads the saved pipeline and does not train.",
        "`src/` — loading, labels, models, evaluation, and the locked-test comparison.",
        "`notebooks/` — executed notebooks 01 to 08, in that order.",
        "`models/` — saved pipelines and the embedding checkpoint.",
        "`results/figures/` and `results/tables/` — saved plots, metrics, and prediction files.",
        "`reports/` — the technical report, phase notes, and the contribution statement.",
        "`tests/` — automated checks.",
        "`data/raw/` — original download. Gitignored. Not included in this package.",
        "`data/processed/` — labelled table written by notebook 01. Gitignored. Not included.",
    ):
        _paragraph(document, item, bullet=True)
    _paragraph(
        document,
        "This package rearranges copies of those files into `Source_Code/`, `Notebooks/`, `Model_Artifacts/`, `Results/`, `Dataset/`, `Documentation/`, and `Presentation/`.",
    )
    document.add_heading("6. Python version and dependencies", 2)
    _paragraph(
        document,
        "Python 3.12.3 was the interpreter used to pin the packages. The application and the frozen pipeline install from `requirements.txt`: pandas 3.0.6, numpy 2.5.4, scipy 1.18.1, scikit-learn 1.9.1, joblib 1.6.0, and streamlit 1.66.0. Notebooks, figures, Excel reading, and the embedding network also need `requirements-research.txt`, which adds matplotlib 3.11.2, seaborn 0.13.2, jupyterlab 4.6.4, openpyxl 3.1.5, and torch 2.14.1+cpu.",
    )
    document.add_heading("7. Installation", 2)
    _code(
        document,
        "python3 -m venv .venv\nsource .venv/bin/activate\npython -m pip install --upgrade pip\npython -m pip install -r requirements.txt\npython -m pip install -r requirements-research.txt",
    )
    _paragraph(document, "On Windows, create the environment with `py -m venv .venv` and activate `.venv\\Scripts\\activate`. The same install commands follow.")
    document.add_heading("8. Dataset placement", 2)
    _paragraph(
        document,
        "The raw extract is not in this package. The project records that the City copy is not redistributed with the code. From the repository root, when `data/raw/` is empty, the documented fetch is:",
    )
    _code(document, "python -m src.fetch_c3")
    _paragraph(
        document,
        "Notebook 01 then writes `data/processed/service_requests_labelled.parquet`. Notebooks 02 to 08 read that labelled file. Do not commit the raw extract. Optional path overrides belong in a gitignored `.env`, copied from `.env.example` only if needed. `Dataset/Data_Folder_Notes.docx` preserves an earlier note about a Google Drive stratified split. That split is not the chronological protocol used by notebooks 01 to 08.",
    )
    document.add_heading("9. Notebook execution order", 2)
    _add_table(
        document,
        [
            ["Step", "Notebook", "Role"],
            ["1", "notebooks/01_data_acquisition_and_audit.ipynb", "Source, raw file, column audit, labels"],
            ["2", "notebooks/02_exploratory_data_analysis.ipynb", "Exploratory figures. Does not train."],
            ["3", "notebooks/03_preprocessing_and_baselines.ipynb", "Leakage-safe features, dummy, logistic regression"],
            ["4", "notebooks/04_traditional_machine_learning.ipynb", "Logistic regression and histogram gradient boosting"],
            ["5", "notebooks/05_neural_networks.ipynb", "Multilayer perceptron and embedding network"],
            ["6", "notebooks/06_evaluation_and_optimisation.ipynb", "Champion tuning and one locked-test score"],
            ["7", "notebooks/07_ablation_and_statistical_testing.ipynb", "Ablations, historical McNemar, Question 6.2"],
            ["8", "notebooks/08_final_model_export.ipynb", "Checks the saved pipeline. Does not refit."],
        ],
    )
    _paragraph(document, "Run each notebook from the repository root with that folder as the working directory. The copies in this package are under `Notebooks/` and keep the `.ipynb` format.")
    document.add_heading("10. Automated tests", 2)
    _code(document, "python -m pytest tests -q")
    _paragraph(document, "A smaller import check is `python -m unittest tests.test_smoke`. In this package the tests are under `Source_Code/tests/`.")
    document.add_heading("11. Streamlit application", 2)
    _code(document, "streamlit run app/app.py")
    _paragraph(
        document,
        "Run that command from the repository root after activating the environment. The process loads `models/frozen_resolution_pipeline.joblib`. It does not fit a model. The form asks for complaint type, suburb, ward, subcouncil, and the request creation date. Work centre is not requested. In this package the same file is `Source_Code/app/app.py` and the pipeline is `Model_Artifacts/frozen_resolution_pipeline.joblib`.",
    )
    document.add_heading("12. Model artifacts and results", 2)
    for item in (
        "`models/frozen_resolution_pipeline.joblib` — the frozen champion used by the application.",
        "`models/embedding_network.pt` and `models/embedding_encoder.joblib` — the approved embedding checkpoint used for Question 6.2. It is not the champion.",
        "`models/dummy_most_frequent.joblib` — the most-frequent baseline artifact.",
        "`results/tables/06_frozen_decision.json` — the validation decision. It records that the locked test was not used for selection.",
        "`results/tables/06_test_predictions.parquet` — locked-test predictions for the champion. Columns are source index, label, prediction, and three probabilities. No notification identifier.",
        "`results/tables/07_mcnemar_test.csv` — historical multilayer-perceptron comparison.",
        "`results/tables/13_question62_record.json` — Question 6.2 embedding comparison.",
        "`results/figures/13_roc_all_models_locked_test.png` and `results/tables/13_roc_all_models_locked_test.csv` — the common ROC figure and class AUCs.",
    ):
        _paragraph(document, item, bullet=True)
    document.add_heading("13. Model selection and reported metrics", 2)
    _paragraph(
        document,
        "Selection used validation macro-F1. The Phase 5 histogram gradient boosting validation macro-F1 is "
        + str(decision["phase5_champion_validation_macro_f1"])
        + ". The embedding network validation macro-F1 is "
        + str(question["comparator_validation_macro_f1"])
        + ". The multilayer perceptron validation macro-F1 recorded on the frozen decision is "
        + str(decision["comparator_phase5_validation_macro_f1"])
        + ". Eight training-period trials were searched. The frozen choice is the untuned refit because its 2025 validation macro-F1 of "
        + str(decision["untuned_validation_macro_f1"])
        + " beat the tuned validation macro-F1 of "
        + str(decision["tuned_validation_macro_f1"])
        + ". The locked test was scored once after that choice. Locked-test accuracy is "
        + locked["accuracy"]
        + ", macro-F1 is "
        + locked["macro_f1"]
        + ", Resolved Slow precision is "
        + slow["precision"]
        + ", Resolved Slow recall is "
        + locked["recall_resolved_slow"]
        + ", Unresolved recall is "
        + locked["recall_unresolved"]
        + ", and macro ROC-AUC is "
        + locked["macro_roc_auc"]
        + ".",
    )
    _paragraph(
        document,
        "Question 6.2 scores the saved embedding checkpoint on the same "
        + f"{question['test_rows']:,}"
        + " rows. Its locked-test macro-F1 in the record is "
        + str(question["locked_test_macro_f1"])
        + ". The continuity-corrected statistic is "
        + str(question["statistic"])
        + " and the p-value is "
        + str(question["p_value"])
        + ". The champion was not changed. The historical perceptron statistic remains 160.6875.",
    )
    document.add_heading("14. Known limitations", 2)
    _paragraph(
        document,
        "Logistic regression, solver "
        + logistic["solver"]
        + " with max_iter "
        + str(logistic["historical_max_iter"])
        + ", raised ConvergenceWarning and did not converge. The comparable validation macro-F1 is "
        + str(logistic["recorded_200_iteration_validation_macro_f1"])
        + ". The 1000-iteration validation macro-F1 was "
        + str(logistic["this_run_validation_macro_f1"])
        + " and was not adopted. The locked test was not scored for that check. Other recorded limits: the seven-day label is a project rule; the extract is frozen; January 2025 is empty in the monthly series; suburb `#` and subcouncil `Not assigned` are stored placeholders; work centre has no assignment time; Resolved Slow precision on the locked test is about 0.210, so most of those flags are the wrong class.",
    )
    document.add_heading("15. Ethical limitations and human review", 2)
    _paragraph(
        document,
        "The report's POPIA section describes withheld identifiers, free text, and coordinates. It does not certify that the processing complies with the Protection of Personal Information Act 4 of 2013. No stakeholder interview was conducted. The brief does not require one. A false negative can leave a slow or unresolved request without earlier attention. The three probabilities are the review aid. A single label is not an instruction to delay, refuse, or prioritise an area.",
    )
    document.add_heading("16. Deployed demonstration", 2)
    paragraph = document.add_paragraph()
    _add_inline(paragraph, "The recorded demonstration address is ")
    _add_hyperlink(paragraph, "https://cape-town-service-classifier.onrender.com", "https://cape-town-service-classifier.onrender.com")
    _add_inline(paragraph, ". This package does not recheck that the live site is currently serving the latest commit. The local command remains `streamlit run app/app.py`.")
    document.add_heading("17. Team", 2)
    _paragraph(
        document,
        "The following identities are the project team. `Documentation/Contribution_Statement.docx` records the supplied division: Thabang Mokoena for problem definition, exploratory analysis, and the Streamlit interface; Mpumelelo Ngozo for the data audit and model development; Selowa Tumiso for preprocessing and the statistical comparisons; Ualwela Netshidzati for the technical report and presentation preparation. Signature cells contain those initials.",
    )
    _add_table(
        document,
        [
            ["Name", "Student number", "Initials", "Email"],
            ["Mpumelelo Ngozo", "218409202", "MSP", "218409202@tut4life.ac.za"],
            ["Selowa Tumiso", "218340172", "TS", "218340172@tut4life.ac.za"],
            ["Ualwela Netshidzati", "220069117", "U", "220069117@tut4life.ac.za"],
            ["Thabang Mokoena", "240668823", "TG", "240668823@tut4life.ac.za"],
        ],
    )
    _paragraph(document, "Group number: not recorded. Complete manually.")
    _paragraph(document, "Lecturer: not recorded. Complete manually.")
    _paragraph(document, "Submission date: not recorded. Complete manually.")
    _paragraph(document, "Institution name beyond the module code IDA117V: not recorded in the repository. Complete manually if the cover requires it.")
    _paragraph(
        document,
        "The technical report PDF in this package is regenerated from `reports/final_report.md`. Its cover lists the four names, initials, student numbers, and email addresses above. The group number, lecturer, and submission date remain “Not recorded.” Appendix A repeats the contribution allocation, initials, and dates.",
    )
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def write_dataset_instructions(destination: Path) -> None:
    document = Document()
    _apply_document_defaults(document, "How to obtain the dataset")
    document.add_heading("How to obtain the dataset", 1)
    _paragraph(
        document,
        "The raw City extract is not in this submission package. It was not present in the working tree when the package was built (`data/raw/` contained only a placeholder). `reports/data_source.md` says the local copy is gitignored and is not redistributed with the code. The City terms are custom. They are not a Creative Commons licence.",
    )
    document.add_heading("Where the project expects the file", 2)
    _paragraph(document, "In the original repository, place the download at `data/raw/`. The filename recorded in the data-source note is `data/raw/c3_notifications.parquet`. Notebook 01 writes the labelled analysis file to `data/processed/service_requests_labelled.parquet`.")
    document.add_heading("How the project obtained it", 2)
    _paragraph(
        document,
        "The dataset page is https://odp-cctegis.opendata.arcgis.com/datasets/90711027d33940b5a06e96ad8a8f7ede_0/explore. The published title is Service Requests 2023 until 06 October 2026 (Deprecated). The owner is the City of Cape Town. On 16 October 2026 the feature layer count was 5,372,466 rows. The README fetch command, from the repository root when `data/raw/` is empty, is:",
    )
    _code(document, "python -m src.fetch_c3")
    _paragraph(document, "The verbatim terms captured on 15 October 2026 are in `Dataset/licence.txt`. The longer source note is `Documentation/Dataset_Source.docx`.")
    document.add_heading("What is not the modelling extract", 2)
    _paragraph(
        document,
        "`data/README.md` also describes a Google Drive folder of stratified `train.parquet`, `val.parquet`, and `test.parquet` files, and a 2,000-row `data/sample/anonymised_sample.csv`. Those files use a different column set and a different split from notebooks 01 to 08. They were not copied into this package as the project dataset. The note itself is preserved as `Dataset/Data_Folder_Notes.docx`.",
    )
    document.add_heading("Dictionary", 2)
    _paragraph(
        document,
        "The delivered frozen pipeline uses `C3_Complaint_Type`, `Suburb`, `Ward`, `Sub_Council`, `created_day_of_week`, `created_month`, and `created_is_weekend`. That list is in `results/tables/06_frozen_decision.json`. `Dataset/Earlier_Modelling_Dictionary.docx` is a faithful copy of `docs/data-dictionary.md`. That older dictionary names `department` and `suburb_group`. It is not the feature list of the frozen champion.",
    )
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def write_package_notes(destination: Path) -> None:
    document = Document()
    _apply_document_defaults(document, "Package notes")
    document.add_heading("Package notes", 1)
    _paragraph(document, "This package was assembled by copying existing repository files and converting Markdown to Word. No model was retrained. No saved metric file was rewritten. Original Markdown files remain in the repository.")
    document.add_heading("Included", 2)
    for item in (
        "Technical report PDF, copied from `reports/final_report.pdf`.",
        "Notebooks 01 to 08 in execution order, still in `.ipynb` format.",
        "Application, source, and tests.",
        "Frozen pipeline, embedding checkpoint, and dummy baseline.",
        "Saved figures and tables, including the locked-test prediction files, the Question 6.2 record, the common ROC figure, ablation tables, and the suburb and ward audits.",
        "Interface screenshots. There is no slide file and no attendance record.",
    ):
        _paragraph(document, item, bullet=True)
    document.add_heading("Not included", 2)
    for item in (
        "The raw City extract and the labelled parquet. They are gitignored, were not in the working tree, and the project says the City copy is not redistributed.",
        "`data/sample/anonymised_sample.csv`. It belongs to an earlier feature set, not notebooks 01 to 08.",
        "Notebooks `q1.ipynb` to `q8.ipynb`. They are an earlier question-template track. `q4` to `q8` contain a single stub cell. They are not the executed sequence.",
        "Duplicate executed-notebook copies under `evidence/reproducibility/`.",
        "`.git`, `.venv`, `__pycache__`, `.pytest_cache`, and `.env`.",
        "A second copy of the technical report in Word. The PDF is the submitted report.",
    ):
        _paragraph(document, item, bullet=True)
    document.add_heading("Items the team must still complete", 2)
    for item in (
        "The contribution statement records the supplied allocation, with initials in the signature cells. Upper rows are dated 5–8 October 2026 and the lower half is dated 12–18 October 2026.",
        "Enter the group number, lecturer, and submission date. They are not recorded. The report cover and the contribution statement name Mpumelelo Ngozo (218409202, 218409202@tut4life.ac.za), Selowa Tumiso (218340172, 218340172@tut4life.ac.za), Ualwela Netshidzati (220069117, 220069117@tut4life.ac.za), and Thabang Mokoena (240668823, 240668823@tut4life.ac.za).",
        "Add presentation or defence evidence only if it exists outside this repository. None was found here.",
        "Redeploy the Streamlit service separately if the live site must match the current application code. This package does not deploy it.",
    ):
        _paragraph(document, item, bullet=True)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def assemble() -> None:
    if PACKAGE.exists():
        shutil.rmtree(PACKAGE)
    PACKAGE.mkdir(parents=True)
    _copy_file(ROOT / "requirements.txt", PACKAGE / "requirements.txt")
    _copy_file(ROOT / "requirements-research.txt", PACKAGE / "requirements-research.txt")
    _copy_file(ROOT / ".python-version", PACKAGE / ".python-version")
    _copy_file(ROOT / "reports" / "final_report.pdf", PACKAGE / "Technical_Report.pdf")
    _copy_file(ROOT / "reports" / "IDA117V_Project_Template.pdf", PACKAGE / "Documentation" / "IDA117V_Project_Template.pdf")
    _copy_file(ROOT / "data" / "licence.txt", PACKAGE / "Dataset" / "licence.txt")
    _copy_file(ROOT / ".env.example", PACKAGE / "Source_Code" / ".env.example")
    _copy_file(ROOT / "render.yaml", PACKAGE / "Source_Code" / "render.yaml")
    _copy_tree(ROOT / "src", PACKAGE / "Source_Code" / "src")
    _copy_tree(ROOT / "app", PACKAGE / "Source_Code" / "app")
    _copy_tree(ROOT / "tests", PACKAGE / "Source_Code" / "tests")
    _copy_tree(ROOT / "scripts", PACKAGE / "Source_Code" / "scripts")
    _copy_tree(ROOT / ".streamlit", PACKAGE / "Source_Code" / ".streamlit")
    _copy_tree(ROOT / "models", PACKAGE / "Model_Artifacts")
    _copy_tree(ROOT / "results" / "figures", PACKAGE / "Results" / "figures")
    _copy_tree(ROOT / "results" / "tables", PACKAGE / "Results" / "tables")
    notebooks = PACKAGE / "Notebooks"
    notebooks.mkdir(parents=True)
    for path in sorted((ROOT / "notebooks").glob("[0-9][0-9]_*.ipynb")):
        _copy_file(path, notebooks / path.name)
    presentation = PACKAGE / "Presentation"
    presentation.mkdir(parents=True)
    for path in sorted((ROOT / "evidence" / "screenshots").glob("*.png")):
        _copy_file(path, presentation / path.name)

    for source_name, destination_name in MARKDOWN_DOCUMENTS:
        source = ROOT / source_name
        if not source.is_file():
            raise FileNotFoundError(source)
        preface = None
        if source_name == "docs/data-dictionary.md":
            preface = [
                "Converted from docs/data-dictionary.md. This older dictionary is not the feature list of the frozen champion. The delivered features are C3_Complaint_Type, Suburb, Ward, Sub_Council, created_day_of_week, created_month, and created_is_weekend.",
            ]
        if source_name == "data/README.md":
            preface = [
                "Converted from data/README.md. The Google Drive files described here are an earlier stratified split. Notebooks 01 to 08 use the chronological split recorded in results/tables/06_frozen_decision.json.",
            ]
        convert_markdown(source, PACKAGE / destination_name, preface=preface)
    write_readme(PACKAGE / "README.docx")
    write_project_overview(PACKAGE / "Documentation" / "Project_Overview.docx")
    write_dataset_instructions(PACKAGE / "Dataset" / "How_to_obtain_the_dataset.docx")
    write_package_notes(PACKAGE / "Documentation" / "Package_Notes.docx")


def _zip_package() -> None:
    if ZIP_PATH.exists():
        ZIP_PATH.unlink()
    with zipfile.ZipFile(ZIP_PATH, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(PACKAGE.rglob("*")):
            if path.is_file():
                archive.write(path, arcname=str(Path(ZIP_ROOT) / path.relative_to(PACKAGE)))


def validate() -> None:
    from pypdf import PdfReader

    expected = [
        PACKAGE / "README.docx",
        PACKAGE / "requirements.txt",
        PACKAGE / "requirements-research.txt",
        PACKAGE / "Technical_Report.pdf",
        PACKAGE / "Documentation" / "Project_Overview.docx",
        PACKAGE / "Documentation" / "Contribution_Statement.docx",
        PACKAGE / "Documentation" / "Presentation_Outline.docx",
        PACKAGE / "Dataset" / "How_to_obtain_the_dataset.docx",
        PACKAGE / "Dataset" / "licence.txt",
        PACKAGE / "Model_Artifacts" / "frozen_resolution_pipeline.joblib",
        PACKAGE / "Results" / "tables" / "06_frozen_decision.json",
        PACKAGE / "Results" / "tables" / "06_test_predictions.parquet",
        PACKAGE / "Results" / "tables" / "13_question62_record.json",
        PACKAGE / "Results" / "tables" / "13_roc_all_models_locked_test.csv",
        PACKAGE / "Results" / "figures" / "13_roc_all_models_locked_test.png",
        PACKAGE / "Results" / "tables" / "07_ablation_summary.csv",
        PACKAGE / "Results" / "tables" / "09_suburb_error_audit.csv",
        PACKAGE / "Results" / "tables" / "09_ward_error_audit.csv",
    ]
    missing = [str(path.relative_to(PACKAGE)) for path in expected if not path.is_file() or path.stat().st_size == 0]
    if missing:
        raise SystemExit("Missing or empty: " + ", ".join(missing))
    notebooks = sorted((PACKAGE / "Notebooks").glob("*.ipynb"))
    if len(notebooks) != 8:
        raise SystemExit(f"Expected 8 notebooks, found {len(notebooks)}")
    for path in notebooks:
        notebook = json.loads(path.read_text(encoding="utf-8"))
        if notebook.get("nbformat") is None or not notebook.get("cells"):
            raise SystemExit(f"Invalid notebook: {path.name}")
    docx_files = list(PACKAGE.rglob("*.docx"))
    for path in docx_files:
        opened = Document(path)
        if not opened.paragraphs and not opened.tables:
            raise SystemExit(f"Empty Word document: {path}")
    reader = PdfReader(str(PACKAGE / "Technical_Report.pdf"))
    if len(reader.pages) < 1:
        raise SystemExit("PDF has no pages")
    text = "\n".join((page.extract_text() or "") for page in reader.pages[:2])
    if "Cape Town" not in text:
        raise SystemExit("PDF text did not contain the project title on the first pages")
    forbidden = []
    for path in PACKAGE.rglob("*"):
        if not path.is_file():
            continue
        name = path.name
        if name in {".env", "secrets.toml"} or path.suffix in {".pyc",}:
            forbidden.append(str(path))
        if "__pycache__" in path.parts or ".git" in path.parts or ".venv" in path.parts:
            forbidden.append(str(path))
    if forbidden:
        raise SystemExit("Forbidden files: " + ", ".join(forbidden))
    with zipfile.ZipFile(ZIP_PATH) as archive:
        bad = archive.testzip()
        if bad:
            raise SystemExit(f"Corrupt zip entry: {bad}")
        names = archive.namelist()
        if any(name.startswith("submission_package/") or name.endswith(".git/") for name in names):
            raise SystemExit("Zip contains an unexpected root or git metadata")
        if not any(name == f"{ZIP_ROOT}/README.docx" for name in names):
            raise SystemExit("Zip is missing the top-level README")
        if not any(name == f"{ZIP_ROOT}/Technical_Report.pdf" for name in names):
            raise SystemExit("Zip is missing the technical report")
    print(f"docx {len(docx_files)}")
    print(f"pdf_pages {len(reader.pages)}")
    print(f"pdf_bytes {(PACKAGE / 'Technical_Report.pdf').stat().st_size}")
    print(f"notebooks {len(notebooks)}")
    print(f"zip_entries {len(names)}")
    print(f"zip_bytes {ZIP_PATH.stat().st_size}")


def main() -> None:
    assemble()
    _zip_package()
    validate()


if __name__ == "__main__":
    main()
